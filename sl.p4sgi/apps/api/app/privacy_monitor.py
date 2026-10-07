"""
PIA compliance monitor (EXPERIMENTAL) for the p4sgi dashboard — 0.4.20.

Tracks the state of the Sierra Leone ULID Privacy Impact Assessment workbook
(PIA_Tool_ULID_SierraLeone_Oct_v6.xlsx) against evidence, instead of against what the
workbook claims. Voluntary alignment with the Data Protection and Right to Access
Information Bill 2025, which is NOT enacted.

Principles
* READ-ONLY toward Google Workspace, DHIS2 and Postgres. It analyses the GAM CSVs the dashboard
  already caches (via the 0.4.19 analysers) and stores only its own small JSON state under
  DATA_DIR/privacy/. No schema change, no learner-level data is ever stored.
* Nothing is "compliant" or "agreed" without evidence. Control state comes from, in order:
  measured (GAM) > evidenced attestation > bare attestation > workbook claim (never scored as observed).
* A DPO contact email is REQUIRED. Until it is set the dashboard shows a blocking state and every
  write endpoint (except setting the DPO and sending feedback) refuses with 409 dpo_required.
* No email is sent by this module. The DPO address is a contact record + a mailto: link in the UI.

Wired from main.py with a guarded include_router. Disable with PRIVACY_MONITOR_ENABLED=0.
"""

from __future__ import annotations

import io
import json
import os
import re
import threading
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Literal

from fastapi import APIRouter, HTTPException, Query, Request, Response
from pydantic import BaseModel, Field

from .privacy_catalogue import AXES, CONTROLS

try:  # 0.4.19 analysers (same package). Degrade, never crash, if absent.
    from . import security_audit as _sa
except Exception:  # noqa: BLE001
    _sa = None  # type: ignore[assignment]

SEED_PATH = Path(__file__).with_name("privacy_seed.json")
MAX_IMPORT_BYTES = 10 * 1024 * 1024
LEGAL_STATUSES = ("Bill", "Assented", "Commenced")

DEFAULT_SETTINGS: dict[str, Any] = {
    "weights": {"verified": 1.0, "evidenced": 0.75, "attested": 0.4, "partial_cap": 0.9,
                "expired_factor": 0.5, "declared": 0.0},
    "met_threshold": 0.98,            # measured fraction at/above which a collector counts as met
    "evidence_valid_days": 365,       # attestation validity when valid_until not given
    "expired_void_days": 365,         # this long past validity the evidence counts 0
    "risk_tolerance": 0.75,           # linked controls must average at least this to credit residual risk
    "consult_sla_days": 60,
    "k_min": 10,                      # suppress unit scores with fewer active accounts than this
    "go_live_threshold": {},          # {axis_no(str): 0-100}; empty = no threshold drawn
    "legal_status": {"status": "Bill", "date": None, "source": None},
    "excluded_controls": {},          # {control_id: reason}
}

_EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)+$")
_PLACEHOLDER_DOMAINS = {"example.com", "example.org", "test.com", "localhost", "domain.com", "email.com"}
_PERSONAL_DOMAINS = {"gmail.com", "yahoo.com", "outlook.com", "hotmail.com", "icloud.com", "proton.me", "protonmail.com"}

# Which workbook risks each control gives evidence for (monitor design, reviewable).
RISK_CONTROLS: dict[str, list[str]] = {
    "R1": ["A04", "A03"], "R2": ["A15", "A14"], "R3": ["A05", "A06"], "R4": ["A19", "A34"],
    "R5": ["A14", "A08", "A13"], "R6": ["A16", "A17"], "R7": ["A18", "A17"], "R8": ["A33"],
    "R9": ["A26", "A25", "A32"], "R10": ["A11", "A12"], "R11": ["A10", "A09", "A39"],
    "R12": ["A17", "A20", "A19"], "R13": ["A22", "A23", "A25", "A29", "A31", "A32"],
    "R14": ["A34", "A12"], "R15": ["A15"], "R16": ["A02"], "R17": ["A17", "A36"], "R18": ["A38", "A40"],
}

# Controls the GAM cache can measure (each leaves a companion evidence gap, hence the 0.9 cap).
COLLECTOR_CONTROLS = ("A20", "A21", "A22", "A23", "A25", "A26", "A32")
COMPUTED_CONTROLS = ("A37", "A40")


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _env_float(name: str, default: float) -> float:
    try:
        return float(os.getenv(name, "") or default)
    except ValueError:
        return default


# ==========================================================================
# Workbook import (pure) — tolerant, and reports defects instead of silently fixing them
# ==========================================================================
def band(score: int | None) -> str | None:
    if score is None:
        return None
    if score >= 16:
        return "Critical"
    if score >= 10:
        return "High"
    if score >= 5:
        return "Medium"
    return "Low"


def declared_value(status: Any) -> float | None:
    """Workbook status -> 0..1 'declared' progress. Unknown vocabulary -> None (reported as a defect)."""
    s = str(status or "").strip().lower()
    if s in ("", "none", "not started", "not yet started", "no", "open", "[select]"):
        return 0.0
    if s in ("scheduled", "planned"):
        return 0.25
    if s in ("in progress", "partial", "partially", "partially implemented", "started", "ongoing"):
        return 0.5
    if s in ("complete", "completed", "implemented", "done", "yes", "compliant", "closed"):
        return 1.0
    return None


def _txt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    return str(v).replace(" ", " ").strip()


def _hdr(s: Any) -> str:
    return re.sub(r"\s+", " ", _txt(s)).lower()


def _find_sheet(wb: Any, start: str) -> Any | None:
    for ws in wb.worksheets:
        if ws.title.strip().lower().startswith(start):
            return ws
    return None


def _rows(ws: Any, header_row: int, first_col: int = 1) -> tuple[list[str], list[list[Any]]]:
    hdr = [_hdr(c.value) for c in ws[header_row]][first_col - 1:]
    out = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        vals = list(r)[first_col - 1:]
        if any(v not in (None, "") for v in vals):
            out.append(vals)
    return hdr, out


def _col(hdr: list[str], *needles: str) -> int | None:
    for i, h in enumerate(hdr):
        if all(n in h for n in needles):
            return i
    return None


def _int(v: Any) -> int | None:
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def parse_workbook(source: bytes | str | Path) -> dict[str, Any]:
    """Parse the PIA workbook into plain dicts. Raises ValueError if it is not a usable PIA workbook."""
    try:
        import openpyxl  # lazy: only needed for import
    except Exception as exc:  # noqa: BLE001
        raise ValueError("openpyxl is not installed in this image; rebuild the image") from exc
    try:
        buf = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
        wb = openpyxl.load_workbook(buf, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"not a readable .xlsx workbook: {exc}") from exc

    defects: list[dict[str, str]] = []

    def defect(sev: str, where: str, msg: str) -> None:
        defects.append({"severity": sev, "where": where, "message": msg})

    # ---- project overview
    project: dict[str, str] = {}
    ws = _find_sheet(wb, "1.")
    if ws:
        for a, b in ws.iter_rows(min_row=1, max_col=2, values_only=True):
            if _txt(a) and _txt(b):
                project[_txt(a)] = _txt(b)

    # ---- tab 3 controls
    controls: list[dict[str, Any]] = []
    ws = _find_sheet(wb, "3.")
    if not ws:
        raise ValueError("tab '3. Country Assessment' not found: this does not look like the PIA workbook")
    hdr_row = next((i for i in range(1, 8) if _hdr(ws.cell(i, 1).value) == "id"), 2)
    hdr, rows = _rows(ws, hdr_row)
    ci = {k: _col(hdr, *n) for k, n in dict(id=("id",), area=("assessment area",), question=("control",),
          status=("status",), evidence=("evidence",), gap=("gap",), action=("corrective",),
          owner=("owner",), target=("target",)).items()}
    seen: set[str] = set()
    for r in rows:
        cid = _txt(r[ci["id"]]) if ci["id"] is not None else ""
        if not re.fullmatch(r"A\d{2}", cid):
            continue
        if cid in seen:
            defect("high", "3. Country Assessment", f"duplicate control id {cid}")
        seen.add(cid)
        g = lambda k: _txt(r[ci[k]]) if ci.get(k) is not None and ci[k] < len(r) else ""  # noqa: E731
        st = g("status")
        dv = declared_value(st)
        if dv is None:
            defect("medium", "3. Country Assessment", f"{cid}: unrecognised status '{st}' (treated as not started)")
            dv = 0.0
        controls.append({"id": cid, "area": g("area"), "question": g("question"), "status": st or "Not started",
                         "declared": dv, "evidence": g("evidence"), "gap": g("gap"), "action": g("action"),
                         "owner": g("owner"), "target": g("target")})
    expected = {f"A{i:02d}" for i in range(1, 41)}
    missing = sorted(expected - seen)
    extra = sorted(seen - expected)
    if missing:
        defect("high", "3. Country Assessment", f"controls missing from the workbook: {', '.join(missing)}")
    if extra:
        defect("medium", "3. Country Assessment", f"controls not in the monitor catalogue (ignored in scoring): {', '.join(extra)}")
    if controls and all(c["declared"] == 0 for c in controls):
        defect("info", "3. Country Assessment", "every control is 'Not started': the workbook declares no progress yet")
    if controls and not any(c["evidence"] for c in controls):
        defect("info", "3. Country Assessment", "the 'Evidence / reference' column is empty for every control")

    # ---- tab 6 actions
    actions: list[dict[str, Any]] = []
    ws = _find_sheet(wb, "6.")
    if ws:
        hdr, rows = _rows(ws, 1)
        ai = {k: _col(hdr, *n) for k, n in dict(id=("action id",), action=("action",), risks=("linked",),
              owner=("owner",), priority=("priority",), target=("target",), status=("status",)).items()}
        ai["action_text"] = next((i for i, h_ in enumerate(hdr) if h_ == "action"), None)
        for r in rows:
            aid = _txt(r[ai["id"]]) if ai["id"] is not None else ""
            if not re.fullmatch(r"ACT-\d+", aid):
                continue
            g = lambda k: _txt(r[ai[k]]) if ai.get(k) is not None and ai[k] < len(r) else ""  # noqa: E731
            risks = [x.strip().upper() for x in re.split(r"[,;/ ]+", g("risks")) if x.strip()]
            tgt = g("target")
            st = g("status")
            if declared_value(st) is None:
                defect("medium", "6. Mitigation Tracker", f"{aid}: unrecognised status '{st}'")
            actions.append({"id": aid, "action": _txt(r[ai["action_text"]]) if ai.get("action_text") is not None else "", "risks": risks,
                            "owner": g("owner"), "priority": g("priority"),
                            "target": "" if (not tgt or "[" in tgt) else tgt, "status": st or "Not started"})
        undated = [a["id"] for a in actions if not a["target"]]
        if undated:
            defect("medium", "6. Mitigation Tracker", f"{len(undated)} of {len(actions)} actions have no target date (placeholder '[Enter date]')")

    # ---- tab 4 risks
    risks: list[dict[str, Any]] = []
    ws = _find_sheet(wb, "4.")
    if ws:
        hdr, rows = _rows(ws, 1)
        def rc(*n: str) -> int | None:
            return _col(hdr, *n)
        ri = dict(id=rc("id"), desc=rc("risk description"), cat=rc("category"), subj=rc("data subjects"),
                  L=rc("likelihood"), I=rc("impact"), iscore=rc("inherent", "score"), irating=rc("inherent", "rating"),
                  mit=rc("mitigation"), rL=rc("residual", "likelihood"), rI=rc("residual", "impact"),
                  rscore=rc("residual", "score"), rrating=rc("residual", "rating"), owner=rc("risk owner"), status=rc("status"))
        for r in rows:
            g = lambda k: r[ri[k]] if ri.get(k) is not None and ri[k] < len(r) else None  # noqa: E731
            rid = _txt(g("id")).upper()
            if not re.fullmatch(r"R\d+", rid):
                continue
            L, I, rL, rI = _int(g("L")), _int(g("I")), _int(g("rL")), _int(g("rI"))
            inh = L * I if L and I else None
            res = rL * rI if rL and rI else None
            if _int(g("iscore")) is not None and inh is not None and _int(g("iscore")) != inh:
                defect("high", "4. Risk Register", f"{rid}: inherent score {g('iscore')} != L x I = {inh} (recomputed)")
            if _txt(g("irating")) and inh is not None and _txt(g("irating")).lower() != (band(inh) or "").lower():
                defect("medium", "4. Risk Register", f"{rid}: inherent rating '{g('irating')}' does not match band for {inh} ({band(inh)})")
            if _int(g("rscore")) is not None and res is not None and _int(g("rscore")) != res:
                defect("high", "4. Risk Register", f"{rid}: residual score {g('rscore')} != L x I = {res} (recomputed)")
            if _txt(g("rrating")) and res is not None and _txt(g("rrating")).lower() != (band(res) or "").lower():
                defect("medium", "4. Risk Register", f"{rid}: residual rating '{g('rrating')}' does not match band for {res} ({band(res)})")
            if inh is not None and res is not None and res > inh:
                defect("high", "4. Risk Register", f"{rid}: residual {res} is higher than inherent {inh}")
            risks.append({"id": rid, "description": _txt(g("desc")), "category": _txt(g("cat")), "subjects": _txt(g("subj")),
                          "L": L, "I": I, "inherent": inh, "inherent_band": band(inh),
                          "mitigation": _txt(g("mit")), "rL": rL, "rI": rI, "residual": res, "residual_band": band(res),
                          "owner": _txt(g("owner")), "status": _txt(g("status"))})
        acts_by_risk: dict[str, list[str]] = defaultdict(list)
        for a in actions:
            for rr in a["risks"]:
                acts_by_risk[rr].append(a["id"])
        rids = {x["id"] for x in risks}
        for a in actions:
            for rr in a["risks"]:
                if rr not in rids:
                    defect("medium", "6. Mitigation Tracker", f"{a['id']} links to unknown risk {rr}")
            if not a["risks"]:
                defect("medium", "6. Mitigation Tracker", f"{a['id']} is linked to no risk")
        for x in risks:
            if not acts_by_risk.get(x["id"]):
                defect("medium", "4. Risk Register", f"{x['id']} has no linked mitigation action")
        credited = [x["id"] for x in risks if x["inherent"] and x["residual"] is not None and x["residual"] < x["inherent"]
                    and not any(declared_value(a["status"]) == 1.0 for a in actions if x["id"] in a["risks"])]
        if credited:
            defect("high", "4. Risk Register",
                   f"{len(credited)} risks show residual < inherent but none of their actions is complete (mitigation credited before it exists): {', '.join(credited)}")

    # ---- tab 5 stakeholders
    stakeholders: list[dict[str, Any]] = []
    ws = _find_sheet(wb, "5.")
    if ws:
        hdr, rows = _rows(ws, 1)
        si = dict(name=_col(hdr, "stakeholder"), role=_col(hdr, "role"), req=_col(hdr, "consultation required"),
                  method=_col(hdr, "method"), status=_col(hdr, "status"), notes=_col(hdr, "notes"), owner=_col(hdr, "action"))
        for n, r in enumerate(rows, 1):
            g = lambda k: _txt(r[si[k]]) if si.get(k) is not None and si[k] < len(r) else ""  # noqa: E731
            if g("name"):
                stakeholders.append({"id": f"S{n:02d}", "name": g("name"), "role": g("role"), "required": g("req"),
                                     "method": g("method"), "status": g("status") or "Not started", "notes": g("notes"), "owner": g("owner")})

    # ---- tab 2 inventory
    inventory: list[dict[str, Any]] = []
    ws = _find_sheet(wb, "2.")
    if ws:
        hdr, rows = _rows(ws, 1)
        ii = dict(n=_col(hdr, "#"), el=_col(hdr, "data element"), subj=_col(hdr, "data subject"), cat=_col(hdr, "category"),
                  src=_col(hdr, "source"), store=_col(hdr, "storage"), ret=_col(hdr, "retention period (suggested"),
                  appr=_col(hdr, "approved"), anon=_col(hdr, "anonymisation"), shared=_col(hdr, "shared with"))
        for r in rows:
            g = lambda k: _txt(r[ii[k]]) if ii.get(k) is not None and ii[k] < len(r) else ""  # noqa: E731
            if g("el"):
                inventory.append({"n": g("n"), "element": g("el"), "subjects": g("subj"), "category": g("cat"), "source": g("src"),
                                  "storage": g("store"), "retention": g("ret"), "retention_approved": g("appr"),
                                  "pseudonymisation": g("anon"), "shared_with": g("shared")})
        unapproved = [x["element"] for x in inventory if not x["retention_approved"] or x["retention_approved"].lower().startswith("not")]
        if inventory and len(unapproved) == len(inventory):
            defect("high", "2. Data Inventory", f"retention is approved for 0 of {len(inventory)} data elements")

    # ---- tab 8 access matrix
    access: list[dict[str, str]] = []
    ws = _find_sheet(wb, "8.")
    if ws:
        hrow = next((i for i in range(1, 8) if _hdr(ws.cell(i, 1).value) == "role"), 3)
        hdr, rows = _rows(ws, hrow)
        for r in rows:
            if _txt(r[0]) and len(r) > 2 and any(_txt(x) for x in r[1:4]):
                access.append({(hdr[i] or f"col{i}"): _txt(v) for i, v in enumerate(r) if i < len(hdr)})

    # ---- tab 10 sign-off
    signoff_roles: list[str] = []
    ws = _find_sheet(wb, "10.")
    filled = 0
    if ws:
        for r in ws.iter_rows(min_row=2, max_row=8, values_only=True):
            if _txt(r[0]) and not _txt(r[0]).lower().startswith("overall"):
                signoff_roles.append(_txt(r[0]))
                filled += sum(1 for v in r[1:4] if _txt(v))
        if filled == 0:
            defect("info", "10. Sign-off", "no signatory has a name, signature or date: agreement is not evidenced")

    # ---- tab 7 legal context
    legal: list[dict[str, str]] = []
    ws = _find_sheet(wb, "7.")
    if ws:
        for a, b in ws.iter_rows(min_row=1, max_col=2, values_only=True):
            if _txt(a) and _txt(b):
                legal.append({"topic": _txt(a), "text": _txt(b)})

    # ---- vocabulary note
    risk_status = sorted({x["status"] for x in risks if x["status"]})
    if len(risk_status) > 3:
        defect("info", "4. Risk Register", f"{len(risk_status)} different risk status wordings in use: {', '.join(risk_status)}")

    return {"project": project, "controls": controls, "risks": risks, "actions": actions, "stakeholders": stakeholders,
            "inventory": inventory, "access": access, "signoff_roles": signoff_roles, "legal": legal,
            "defects": defects, "imported_at": _now().isoformat(timespec="seconds")}


def load_seed() -> dict[str, Any]:
    try:
        return json.loads(SEED_PATH.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {"project": {}, "controls": [], "risks": [], "actions": [], "stakeholders": [], "inventory": [], "access": [],
                "signoff_roles": [], "legal": [], "defects": [{"severity": "high", "where": "seed", "message": "bundled workbook seed missing: import the PIA workbook"}],
                "imported_at": None}


# ==========================================================================
# Settings / DPO validation (pure)
# ==========================================================================
def validate_dpo_email(raw: str) -> tuple[str, list[str]]:
    email = (raw or "").strip().lower()
    if not email or len(email) > 254 or not _EMAIL_RE.match(email):
        raise ValueError("enter a valid email address for the Data Protection Officer / focal point")
    dom = email.rsplit("@", 1)[1]
    if dom in _PLACEHOLDER_DOMAINS or email.split("@")[0] in ("test", "dpo", "example", "name", "email", "none", "noreply", "no-reply"):
        raise ValueError("that looks like a placeholder address; enter the real DPO / focal-point mailbox")
    warns = []
    if dom in _PERSONAL_DOMAINS:
        warns.append("personal_mailbox: a role mailbox on an institutional domain is recommended so the contact survives staff changes")
    return email, warns


def merged_settings(stored: dict[str, Any] | None) -> dict[str, Any]:
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    for k, v in (stored or {}).items():
        if k == "weights" and isinstance(v, dict):
            s["weights"].update({a: b for a, b in v.items() if a in s["weights"]})
        elif k in s:
            s[k] = v
    return s


# ==========================================================================
# Evidence-based control state (pure)
# ==========================================================================
def _parse_date(v: Any) -> datetime | None:
    if not v:
        return None
    try:
        d = datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    except ValueError:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def attestation_result(att: dict[str, Any], w: dict[str, float], cfg: dict[str, Any], now: datetime) -> dict[str, Any]:
    stt = att.get("state")
    pct = float(att.get("pct") if att.get("pct") is not None else (1.0 if stt == "met" else 0.0))
    tier = "evidenced" if (att.get("evidence_ref") or "").strip() else "attested"
    if stt == "gap":
        score, state = 0.0, "gap"
    else:
        base = 1.0 if stt == "met" else min(max(pct, 0.0), w["partial_cap"])
        score = base * w[tier]
        state = tier if stt == "met" else "partial"
    until = _parse_date(att.get("valid_until")) or ((_parse_date(att.get("at")) or now) + timedelta(days=int(cfg["evidence_valid_days"])))
    expired = until < now
    void = expired and (now - until).days > int(cfg["expired_void_days"])
    if void:
        score = 0.0
    elif expired:
        score *= w["expired_factor"]
    return {"state": state, "score": round(score, 4), "tier": tier, "expired": expired, "void": void,
            "valid_until": until.date().isoformat()}


def collector_result(frac: float, cap: float, w: dict[str, float], cfg: dict[str, Any]) -> tuple[str, float]:
    if frac <= 0:
        return "gap", 0.0
    if frac >= cfg["met_threshold"] and cap >= 1.0:
        return "verified", w["verified"]
    return "partial", round(min(frac, cap), 4)


CATALOGUE = {c[0]: c for c in CONTROLS}


def control_states(wb: dict[str, Any], measured: dict[str, dict[str, Any]], attests: dict[str, dict[str, Any]],
                   cfg: dict[str, Any], unit: str | None, dpo_email: str | None, now: datetime) -> dict[str, dict[str, Any]]:
    """Per-control evidence state for one scope (unit=None means national)."""
    w = cfg["weights"]
    wb_by_id = {c["id"]: c for c in wb.get("controls", [])}
    out: dict[str, dict[str, Any]] = {}
    for cid, axis, ev, auto, bill, measure in CONTROLS:
        wc = wb_by_id.get(cid, {})
        row = {"id": cid, "axis": axis, "evidence": ev, "automation": auto, "bill": bill, "measure": measure,
               "question": wc.get("question") or measure, "area": wc.get("area", ""),
               "declared_status": wc.get("status", "Not started"), "declared": wc.get("declared", 0.0),
               "state": "no_data", "score": None, "basis": "none", "detail": "", "conflict": False,
               "expired": False, "attestation": None, "excluded": cid in cfg["excluded_controls"]}
        if cid in cfg["excluded_controls"]:
            row["detail"] = "excluded by configuration: " + str(cfg["excluded_controls"][cid])
            out[cid] = row
            continue
        att = attests.get(f"{cid}|{unit or ''}")
        a_res = attestation_result(att, w, cfg, now) if att else None
        if att:
            row["attestation"] = {k: att.get(k) for k in ("state", "pct", "note", "evidence_ref", "valid_until", "by", "at")}
        m = measured.get(cid)
        if m is not None:
            lifted = bool(a_res and att.get("state") == "met" and a_res["tier"] == "evidenced" and not a_res["expired"])
            cap = 1.0 if lifted else w["partial_cap"]
            state, score = collector_result(m["frac"], cap, w, cfg)
            row.update(state=state, score=score, basis="measured", detail=m["detail"], measured=m)
            if att and att.get("state") == "met" and m["frac"] < 0.9:
                row["conflict"] = True
                row["detail"] += "  | CONFLICT: attested 'met' but measurement is below 90%"
            elif not lifted and m["frac"] >= cfg["met_threshold"]:
                row["detail"] += f"  | capped at {w['partial_cap']:.0%}: remaining evidence ({ev}) needs an evidenced attestation"
        elif a_res:
            row.update(state=a_res["state"], score=a_res["score"], basis="attestation", expired=a_res["expired"],
                       detail=("evidence expired " + a_res["valid_until"]) if a_res["expired"] else "attested by " + str(att.get("by")))
            if a_res["void"]:
                row["state"], row["detail"] = "no_data", "attestation lapsed more than a year ago: counts 0"
                row["score"] = 0.0
        elif cid == "A02" and unit is None and dpo_email:
            syn = {"state": "partial", "pct": 0.5, "evidence_ref": "", "at": None}
            r = attestation_result(syn, w, cfg, now)
            row.update(state="partial", score=r["score"], basis="dpo_setting",
                       detail=f"DPO contact registered in the monitor ({dpo_email}); a signed designation letter is not evidenced yet")
        out[cid] = row
    return out


def finalize_computed(states: dict[str, dict[str, Any]], gate_open: bool, now: datetime) -> None:
    """A37 (evidence production) and A40 (gate) are computed, not collected."""
    scored = [s for s in states.values() if s["id"] not in COMPUTED_CONTROLS and not s["excluded"]]
    fresh = [s for s in scored if s["state"] in ("verified", "evidenced") and not s["expired"]]
    if "A37" in states and not states["A37"]["excluded"]:
        f = len(fresh) / len(scored) if scored else 0.0
        states["A37"].update(state="verified" if f >= 0.98 else ("partial" if f > 0 else "gap"), score=round(min(f, 1.0), 4),
                             basis="computed", detail=f"{len(fresh)} of {len(scored)} controls have fresh verified/evidenced support")
    if "A40" in states and not states["A40"]["excluded"]:
        states["A40"].update(state="verified" if gate_open else "gap", score=1.0 if gate_open else 0.0, basis="computed",
                             detail="go-live gate is OPEN" if gate_open else "go-live gate is BLOCKED (see Gate)")


def axis_scores(states: dict[str, dict[str, Any]], cfg: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for n, name in AXES:
        ids = [c for c, s in states.items() if s["axis"] == n and not s["excluded"]]
        sc = [states[c]["score"] for c in ids if states[c]["score"] is not None]
        dec = [states[c]["declared"] for c in ids]
        out.append({
            "axis": n, "name": name, "controls": len(ids), "scored": len(sc),
            "coverage": round(len(sc) / len(ids), 3) if ids else 0.0,
            "observed": round(100 * sum(sc) / len(sc), 1) if sc else None,   # over measurable controls only
            "assured": round(100 * sum(sc) / len(ids), 1) if ids else None,  # unmeasured counts as 0
            "declared": round(100 * sum(dec) / len(dec), 1) if dec else None,
            "ideal": 100, "threshold": cfg["go_live_threshold"].get(str(n)),
        })
    return out


# ==========================================================================
# Risks, gate, sign-off, tiles (pure)
# ==========================================================================
def _action_status(a: dict[str, Any], overlay: dict[str, Any]) -> str:
    return (overlay.get(a["id"]) or {}).get("status") or a["status"]


def risk_view(wb: dict[str, Any], states: dict[str, dict[str, Any]], overlay: dict[str, Any],
              acceptances: dict[str, Any], cfg: dict[str, Any], now: datetime) -> list[dict[str, Any]]:
    out = []
    actions = wb.get("actions", [])
    for r in wb.get("risks", []):
        linked_actions = [a for a in actions if r["id"] in a["risks"]]
        astat = {a["id"]: _action_status(a, overlay) for a in linked_actions}
        all_done = bool(linked_actions) and all(declared_value(s) == 1.0 for s in astat.values())
        ctl_ids = [c for c in RISK_CONTROLS.get(r["id"], []) if c in states and not states[c]["excluded"]]
        sc = [states[c]["score"] for c in ctl_ids]
        scored = [x for x in sc if x is not None]
        mean = (sum(scored) / len(ctl_ids)) if ctl_ids else None  # unmeasured counts as 0 (strict)
        in_tol = mean is not None and len(scored) == len(ctl_ids) and mean >= cfg["risk_tolerance"]
        why = []
        if not linked_actions:
            why.append("no linked action")
        elif not all_done:
            why.append("linked action(s) not complete: " + ", ".join(f"{k}={v}" for k, v in astat.items() if declared_value(v) != 1.0))
        if not in_tol:
            why.append("evidence for linked controls below tolerance" + (f" (mean {mean:.2f}, {len(scored)}/{len(ctl_ids)} controls measured)" if mean is not None else ""))
        credited = all_done and in_tol
        adj = r["residual"] if credited and r["residual"] is not None else r["inherent"]
        acc = acceptances.get(r["id"])
        accepted = False
        if acc:
            rb = _parse_date(acc.get("review_by"))
            accepted = not (rb and rb < now)
        out.append({**r, "actions": astat, "controls": ctl_ids, "control_mean": None if mean is None else round(mean, 3),
                    "controls_measured": len(scored), "credited": credited, "adjusted": adj, "adjusted_band": band(adj),
                    "why_not_credited": [] if credited else why, "acceptance": acc, "accepted": accepted,
                    "stated_status": r["status"],
                    "flag": ("accepted without an acceptance record" if str(r["status"]).lower().startswith("accepted") and not acc else None)})
    return out


SIGNOFF_STATES = ("Draft", "In review", "Signed", "Decision recorded")


def signoff_view(wb: dict[str, Any], overlay: dict[str, Any]) -> dict[str, Any]:
    roles = wb.get("signoff_roles") or []
    rows = []
    signed = 0
    for r in roles:
        o = overlay.get("roles", {}).get(r) or {}
        ok = bool(o.get("name") and o.get("date"))
        signed += ok
        rows.append({"role": r, "name": o.get("name", ""), "date": o.get("date", ""), "decision": o.get("decision", ""),
                     "evidence_ref": o.get("evidence_ref", ""), "signed": ok, "by": o.get("by")})
    decision = overlay.get("decision") or {}
    if decision.get("value") and roles and signed == len(roles):
        st = "Decision recorded"
    elif roles and signed == len(roles):
        st = "Signed"
    elif signed or any(r["name"] for r in rows):
        st = "In review"
    else:
        st = "Draft"
    return {"state": st, "signed": signed, "total": len(roles), "rows": rows, "decision": decision,
            "agreement": "evidenced" if st in ("Signed", "Decision recorded") else "not evidenced",
            "conditions": overlay.get("conditions", "")}


def gate_view(risks: list[dict[str, Any]], signoff: dict[str, Any], dpo_email: str | None) -> dict[str, Any]:
    blockers, conditions = [], []
    if not dpo_email:
        blockers.append({"kind": "dpo_missing", "ref": "ACT-01 / R16", "message": "No Data Protection Officer / focal-point contact has been entered"})
    for r in risks:
        if r["accepted"] or r["adjusted_band"] not in ("Critical", "High"):
            continue
        item = {"kind": f"risk_{r['adjusted_band'].lower()}", "ref": r["id"], "score": r["adjusted"],
                "message": f"{r['id']} is {r['adjusted_band']} on evidence ({r['adjusted']}); " + ("; ".join(r["why_not_credited"]) or "no credit available")}
        (blockers if r["adjusted_band"] == "Critical" else conditions).append(item)
    if signoff["state"] not in ("Signed", "Decision recorded"):
        blockers.append({"kind": "signoff", "ref": "tab 10", "message": f"Sign-off is '{signoff['state']}' ({signoff['signed']}/{signoff['total']} signatories): agreement is not evidenced"})
    status = "BLOCKED" if blockers else ("CONDITIONAL" if conditions else "OPEN")
    return {"status": status, "blockers": blockers, "conditions": conditions,
            "rule": "Critical (not accepted) or missing DPO or incomplete sign-off blocks approval/go-live; High (not accepted) must be closed before go-live of the affected component."}


def triangulation(national: list[dict[str, Any]], states: dict[str, dict[str, Any]], unit_axes: dict[str, list[dict[str, Any]]],
                  unit_overall: dict[str, float | None]) -> dict[str, Any]:
    def mean(vals: list[float]) -> float | None:
        v = [x for x in vals if x is not None]
        return round(sum(v) / len(v), 1) if v else None

    decl = mean([a["declared"] for a in national])
    obs = mean([a["observed"] for a in national])
    assu = mean([a["assured"] for a in national])
    total = [s for s in states.values() if not s["excluded"]]
    scored = [s for s in total if s["score"] is not None]
    att = [s for s in total if s["basis"] == "attestation"]
    fresh = [s for s in att if not s["expired"]]
    # equity spread across units, per axis (>=3 units observed)
    spreads = []
    for i, (n, name) in enumerate(AXES):
        vals = [ax[i]["observed"] for ax in unit_axes.values() if ax[i]["observed"] is not None]
        if len(vals) >= 3:
            spreads.append({"axis": n, "name": name, "spread": round(max(vals) - min(vals), 1), "units": len(vals)})
    spreads.sort(key=lambda x: -x["spread"])
    ov = {u: v for u, v in unit_overall.items() if v is not None}
    worst = min(ov.items(), key=lambda kv: kv[1]) if ov else None
    contradictions = [s["id"] for s in total if s["declared"] >= 1.0 and s["state"] in ("gap", "no_data")]
    return {
        "ambition_gap": None if decl is None else round(100 - decl, 1),
        "assurance_gap": None if decl is None or assu is None else round(decl - assu, 1),
        "assurance_gap_note": "declared minus evidence-assured (unmeasured counts as 0); positive = claims not borne out; negative = measured evidence exceeds what the workbook claims",
        "declared_but_unevidenced": contradictions,
        "equity_spread": spreads[0] if spreads else None, "equity_by_axis": spreads,
        "tail_risk": None if not worst else {"unit": worst[0], "score": worst[1], "units_scored": len(ov)},
        "evidence_freshness": None if not att else round(len(fresh) / len(att), 3),
        "coverage": round(len(scored) / len(total), 3) if total else 0.0,
        "coverage_counts": {"scored": len(scored), "total": len(total)},
        "national_observed": obs, "national_assured": assu, "national_declared": decl,
    }


# ==========================================================================
# Collectors over cached GAM CSVs (pure)
# ==========================================================================
def _domain_of(email: str) -> str:
    return email.rsplit("@", 1)[1].lower() if "@" in email else ""


def compute_collectors(users: list[dict[str, str]], logins: list[dict[str, str]], admins: list[dict[str, str]],
                       dev_rows: list[dict[str, Any]], log_counts: dict[str, int], row_email: Callable[..., str],
                       now: datetime) -> tuple[dict[str, dict[str, Any]], int]:
    """Return ({control_id: {frac, detail, n}}, n_active_accounts). Aggregates only; no person-level output."""
    truthy = _sa.truthy if _sa else (lambda v: None)
    measured: dict[str, dict[str, Any]] = {}
    active = [u for u in users if (truthy(u.get("suspended")) is not True)]
    n_active = len(active)
    by_email = {row_email(u, "primaryEmail"): u for u in active}

    known = [u for u in active if truthy(u.get("isEnrolledIn2Sv")) is not None]
    if known:
        enr = sum(1 for u in known if truthy(u.get("isEnrolledIn2Sv")) is True)
        measured["A22"] = {"frac": enr / len(known), "n": len(known),
                           "detail": f"2-step verification enrolled for {enr} of {len(known)} active accounts"}
    adm = {row_email(a, "primaryEmail", "email", "user", "User") for a in admins} - {""}
    adm_users = [by_email[e] for e in adm if e in by_email and truthy(by_email[e].get("isEnrolledIn2Sv")) is not None]
    if adm_users:
        ok = sum(1 for u in adm_users if truthy(u.get("isEnrolledIn2Sv")) is True)
        measured["A20"] = {"frac": ok / len(adm_users), "n": len(adm_users),
                           "detail": f"{ok} of {len(adm_users)} administrator accounts have 2-step verification"}
    if _sa and active:
        res = _sa.analyse_users(users, logins, admins, lambda e: True, row_email, now)
        dormant = sum(1 for r in res["rows"] if "INACTIVE" in r["flags"] or "NEVER_LOGGED_IN" in r["flags"])
        measured["A21"] = {"frac": 1 - dormant / n_active, "n": n_active,
                           "detail": f"{dormant} of {n_active} active accounts are dormant (inactive or never logged in); a quarterly review record is still needed"}
    if dev_rows:
        n = len(dev_rows)
        comp = sum(1 for d in dev_rows if d["compliance_status"] == "COMPLIANT")
        enc = sum(1 for d in dev_rows if "UNENCRYPTED" in d["flags"])
        pat = sum(1 for d in dev_rows if "OUTDATED_PATCH" in d["flags"])
        measured["A26"] = {"frac": comp / n, "n": n, "detail": f"{comp} of {n} managed devices are fully compliant"}
        measured["A25"] = {"frac": 1 - enc / n, "n": n, "detail": f"{n - enc} of {n} devices report disk encryption (hosting/backup encryption needs attestation)"}
        measured["A32"] = {"frac": 1 - pat / n, "n": n, "detail": f"{n - pat} of {n} devices have a security patch newer than the limit (server/dependency scans need attestation)"}
    srcs = ("login_activity", "token_activity", "drive_activity")
    present = sum(1 for k in srcs if (log_counts.get(k) or 0) > 0)
    if any(log_counts.get(k) is not None for k in srcs):  # no cached audit report at all = no claim, not a gap
            measured["A23"] = {"frac": present / 3, "n": 3,
                           "detail": f"audit data present for {present} of 3 sources (login, token, drive); DHIS2 audit log not observed here"}
    return measured, n_active


# ==========================================================================
# Persistent state (own JSON files; never touches Postgres)
# ==========================================================================
class Store:
    def __init__(self, base: Path):
        self.base = base
        self.lock = threading.Lock()

    def _p(self, name: str) -> Path:
        return self.base / name

    def read_state(self) -> dict[str, Any]:
        try:
            return json.loads(self._p("state.json").read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            return {}

    def mutate(self, fn: Callable[[dict[str, Any]], Any]) -> Any:
        with self.lock:
            st = self.read_state()
            res = fn(st)
            self.base.mkdir(parents=True, exist_ok=True)
            tmp = self._p("state.json.tmp")
            tmp.write_text(json.dumps(st, indent=1, sort_keys=True), encoding="utf-8")
            os.replace(tmp, self._p("state.json"))
            return res

    def append(self, name: str, rec: dict[str, Any]) -> None:
        with self.lock:
            self.base.mkdir(parents=True, exist_ok=True)
            with self._p(name).open("a", encoding="utf-8") as f:
                f.write(json.dumps(rec, sort_keys=True) + "\n")

    def read_lines(self, name: str, limit: int = 500) -> list[dict[str, Any]]:
        p = self._p(name)
        if not p.exists():
            return []
        out = []
        for line in p.read_text(encoding="utf-8").splitlines()[-limit:]:
            try:
                out.append(json.loads(line))
            except ValueError:
                continue
        return out

    def workbook(self) -> dict[str, Any]:
        try:
            return json.loads(self._p("workbook.json").read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            return {}

    def save_workbook(self, wb: dict[str, Any]) -> str:
        with self.lock:
            self.base.mkdir(parents=True, exist_ok=True)
            ver = _now().strftime("%Y%m%dT%H%M%SZ")
            (self.base / "workbooks").mkdir(exist_ok=True)
            data = json.dumps(wb, indent=1)
            (self.base / "workbooks" / f"{ver}.json").write_text(data, encoding="utf-8")
            tmp = self._p("workbook.json.tmp")
            tmp.write_text(data, encoding="utf-8")
            os.replace(tmp, self._p("workbook.json"))
            return ver


# ==========================================================================
# Request models
# ==========================================================================
class DpoIn(BaseModel):
    email: str = Field(..., max_length=254)
    name: str | None = Field(None, max_length=120)
    role: str | None = Field(None, max_length=120)
    designation_ref: str | None = Field(None, max_length=500, description="link/vault path of the signed designation letter (optional)")


class SettingsIn(BaseModel):
    weights: dict[str, float] | None = None
    met_threshold: float | None = Field(None, ge=0.5, le=1.0)
    evidence_valid_days: int | None = Field(None, ge=30, le=1825)
    expired_void_days: int | None = Field(None, ge=30, le=1825)
    risk_tolerance: float | None = Field(None, ge=0.1, le=1.0)
    consult_sla_days: int | None = Field(None, ge=7, le=730)
    k_min: int | None = Field(None, ge=1, le=1000)
    go_live_threshold: dict[str, float] | None = None
    legal_status: dict[str, Any] | None = None
    excluded_controls: dict[str, str] | None = None


class AttestationIn(BaseModel):
    control_id: str = Field(..., pattern=r"^A\d{2}$")
    unit: str | None = Field(None, max_length=120, description="email domain; omit for national")
    state: Literal["met", "partial", "gap"]
    pct: float | None = Field(None, ge=0, le=1)
    note: str = Field("", max_length=2000)
    evidence_ref: str = Field("", max_length=500, description="link / Vault matter / Drive file id of the supporting document")
    valid_until: str | None = Field(None, max_length=32)


class TrackerIn(BaseModel):
    kind: Literal["action", "stakeholder"]
    id: str = Field(..., max_length=20)
    status: str | None = Field(None, max_length=60)
    owner: str | None = Field(None, max_length=200)
    target_date: str | None = Field(None, max_length=32)
    evidence_ref: str | None = Field(None, max_length=500)
    note: str | None = Field(None, max_length=1000)


class SignoffIn(BaseModel):
    role: str | None = Field(None, max_length=200)
    name: str | None = Field(None, max_length=160)
    date: str | None = Field(None, max_length=32)
    decision: str | None = Field(None, max_length=200)
    evidence_ref: str | None = Field(None, max_length=500)
    overall_decision: Literal["Proceed", "Proceed with conditions", "Do not proceed"] | None = None
    conditions: str | None = Field(None, max_length=4000)


class AcceptanceIn(BaseModel):
    risk_id: str = Field(..., pattern=r"^R\d+$")
    note: str = Field(..., min_length=5, max_length=1000)
    review_by: str = Field(..., max_length=32)
    revoke: bool = False


class FeedbackIn(BaseModel):
    category: Literal["bug", "idea", "wrong_control", "threshold", "usability", "other"] = "other"
    control_id: str | None = Field(None, pattern=r"^(A\d{2}|R\d+|ACT-\d+)?$")
    rating: int | None = Field(None, ge=1, le=5)
    text: str = Field(..., min_length=3, max_length=2000)


# ==========================================================================
# Router
# ==========================================================================
@dataclass
class Helpers:
    insight_scope: Callable[..., tuple[set[str] | None, dict[str, Any]]]
    load_sources: Callable[[list[str], set[str] | None], tuple[dict[str, list[dict[str, str]]], dict[str, Any]]]
    build_device_rows: Callable[[dict[str, list[dict[str, str]]], set[str] | None], list[dict[str, Any]]]
    row_email: Callable[..., str]
    in_scope: Callable[[str, set[str] | None], bool]
    request_scope: Callable[[Request], dict[str, Any]]
    is_superadmin: Callable[[str | None], bool]
    data_dir: Path
    app_version: str


SOURCES = ["users_full", "login_activity", "admins", "token_activity", "drive_activity",
           "devices_ci", "devices_mobile", "devices_cros"]


def build_router(h: Helpers) -> APIRouter:
    router = APIRouter(prefix="/api/v1/privacy", tags=["privacy-monitor"])
    store = Store(h.data_dir / "privacy")

    # ---------------- state helpers
    def _workbook() -> tuple[dict[str, Any], str]:
        wb = store.workbook()
        if wb.get("controls"):
            return wb, "imported"
        return load_seed(), "bundled_seed"

    def _cfg(st: dict[str, Any]) -> dict[str, Any]:
        cfg = merged_settings(st.get("settings"))
        cfg["k_min"] = int(cfg["k_min"] or _env_float("PRIV_K_MIN", 10))
        return cfg

    def _dpo(st: dict[str, Any]) -> dict[str, Any]:
        d = st.get("dpo") or {}
        return {"configured": bool(d.get("email")), **{k: d.get(k) for k in ("email", "name", "role", "designation_ref", "set_by", "set_at")}}

    def _actor(request: Request) -> tuple[dict[str, Any], str]:
        sc = h.request_scope(request)
        return sc, str(sc.get("real_email") or sc.get("email") or "local")

    def _need_super(request: Request) -> tuple[dict[str, Any], str]:
        sc, who = _actor(request)
        if not sc.get("superadmin"):
            raise HTTPException(403, {"error": "superadmin_required"})
        return sc, who

    def _need_dpo(st: dict[str, Any]) -> None:
        if not _dpo(st)["configured"]:
            raise HTTPException(409, {"error": "dpo_required",
                                      "message": "Enter the Data Protection Officer / focal-point email first (PUT /api/v1/privacy/dpo)."})

    def _audit(who: str, action: str, detail: dict[str, Any]) -> None:
        store.append("audit.jsonl", {"at": _now().isoformat(timespec="seconds"), "by": who, "action": action, "detail": detail})

    # ---------------- model assembly
    def _assemble(request: Request, domain: list[str] | None, domains: str | None) -> dict[str, Any]:
        now = _now()
        st = store.read_state()
        cfg = _cfg(st)
        dpo = _dpo(st)
        wb, wb_src = _workbook()
        allowed, _ = h.insight_scope(request, domain, domains)
        data, info = h.load_sources(SOURCES, allowed)
        sources = {k: {"status": v.get("status"), "cached_at": v.get("cached_at"), "rows": v.get("rows_raw")} for k, v in info.items()}
        re_ = h.row_email

        def dom_rows(rows: list[dict[str, str]], *keys: str) -> dict[str, list[dict[str, str]]]:
            g: dict[str, list[dict[str, str]]] = defaultdict(list)
            for r in rows:
                e = re_(r, *keys)
                if e and h.in_scope(e, allowed):
                    g[_domain_of(e)].append(r)
            return g

        users_g = dom_rows(data.get("users_full", []), "primaryEmail")
        logins_g = dom_rows(data.get("login_activity", []), "actor.email")
        admins_g = dom_rows(data.get("admins", []), "primaryEmail", "email", "user", "User")
        tok_g = dom_rows(data.get("token_activity", []), "actor.email")
        drv_g = dom_rows(data.get("drive_activity", []), "actor.email")
        dev_rows: list[dict[str, Any]] = []
        if _sa:
            try:
                dev_rows = _sa.analyse_devices(h.build_device_rows(data, allowed), now)["rows"]
            except Exception:  # noqa: BLE001
                dev_rows = []
        dev_g: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for d in dev_rows:
            dm = (d.get("domain") or _domain_of(d.get("email") or "")).lower()
            if dm:
                dev_g[dm].append(d)

        def flat(g: dict[str, list[Any]]) -> list[Any]:
            return [x for v in g.values() for x in v]

        attests = st.get("attestations") or {}
        cached = {k: (info.get(k) or {}).get("status") == "ok" for k in ("login_activity", "token_activity", "drive_activity")}

        def cnt(k: str, g: dict[str, list[Any]], u: str | None = None) -> int | None:
            if not cached[k]:
                return None
            return len(g.get(u, [])) if u else sum(len(v) for v in g.values())

        nat_meas, nat_n = compute_collectors(
            flat(users_g), flat(logins_g), flat(admins_g), dev_rows,
            {"login_activity": cnt("login_activity", logins_g), "token_activity": cnt("token_activity", tok_g),
             "drive_activity": cnt("drive_activity", drv_g)}, re_, now)

        sign = signoff_view(wb, st.get("signoff") or {})
        # two passes: risks/gate depend on control states, A37/A40 depend on the gate.
        nat = control_states(wb, nat_meas, attests, cfg, None, dpo.get("email"), now)
        risks = risk_view(wb, nat, st.get("tracker", {}).get("action", {}), st.get("acceptances") or {}, cfg, now)
        gate = gate_view(risks, sign, dpo.get("email"))
        finalize_computed(nat, gate["status"] == "OPEN", now)
        risks = risk_view(wb, nat, st.get("tracker", {}).get("action", {}), st.get("acceptances") or {}, cfg, now)
        gate = gate_view(risks, sign, dpo.get("email"))
        nat_axes = axis_scores(nat, cfg)

        units: dict[str, dict[str, Any]] = {}
        unit_axes: dict[str, list[dict[str, Any]]] = {}
        unit_overall: dict[str, float | None] = {}
        for u in sorted(users_g):
            meas, n = compute_collectors(
                users_g[u], logins_g.get(u, []), admins_g.get(u, []), dev_g.get(u, []),
                {"login_activity": cnt("login_activity", logins_g, u), "token_activity": cnt("token_activity", tok_g, u),
                 "drive_activity": cnt("drive_activity", drv_g, u)}, re_, now)
            suppressed = n < cfg["k_min"]
            stt = control_states(wb, meas, attests, cfg, u, None, now)
            ax = axis_scores(stt, cfg)
            obs = [a["observed"] for a in ax if a["observed"] is not None]
            overall = None if suppressed or not obs else round(sum(obs) / len(obs), 1)
            units[u] = {"unit": u, "accounts": None if suppressed else n, "suppressed": suppressed, "overall": overall,
                        "scored_controls": sum(1 for s in stt.values() if s["score"] is not None)}
            unit_axes[u] = ax if not suppressed else [{**a, "observed": None, "assured": None} for a in ax]
            unit_overall[u] = overall
            units[u]["_states"] = stt
        tri = triangulation(nat_axes, nat, unit_axes, unit_overall)
        return {"now": now, "cfg": cfg, "dpo": dpo, "wb": wb, "wb_src": wb_src, "st": st, "allowed": allowed,
                "sources": sources, "nat": nat, "nat_axes": nat_axes, "risks": risks, "gate": gate, "sign": sign,
                "units": units, "unit_axes": unit_axes, "tri": tri, "n_accounts": nat_n}

    def _public_units(m: dict[str, Any]) -> list[dict[str, Any]]:
        return [{k: v for k, v in u.items() if not k.startswith("_")} for u in m["units"].values()]

    def _scope_label(m: dict[str, Any]) -> list[str]:
        return sorted(m["allowed"]) if m["allowed"] else ["all"]

    # ---------------- endpoints: read
    @router.get("/status")
    def status(request: Request) -> dict:
        sc = h.request_scope(request)
        st = store.read_state()
        wb, src = _workbook()
        return {
            "module": "privacy_monitor", "version": h.app_version, "experimental": True, "mode": "read_only_toward_workspace",
            "superadmin": bool(sc.get("superadmin")), "dpo": _dpo(st), "workbook_source": src,
            "workbook_imported_at": wb.get("imported_at"), "defects": len(wb.get("defects", [])),
            "legal": {"framework": "Data Protection and Right to Access Information Bill 2025 — NOT enacted; voluntary alignment only",
                      **merged_settings(st.get("settings"))["legal_status"]},
            "counts": {"controls": len(CONTROLS), "measurable_from_gam": len(COLLECTOR_CONTROLS), "computed": len(COMPUTED_CONTROLS)},
            "limits": [
                "Only 7 of 40 controls can be measured from the cached GAM reports (A20, A21, A22, A23, A25, A26, A32), and each is capped at 90% until the companion evidence is attested.",
                "DHIS2 and Google Vault collectors are not built: those controls need an evidenced attestation (link to the document) until they are.",
                "Unit (district/school) polygons are per email domain; the domain-to-district/school mapping is not in the cached data.",
                "Nothing is marked compliant or agreed without evidence. The workbook's own status is shown as 'declared' and never scored as observed.",
                "No email is sent. The DPO address is a contact record with a mailto link.",
            ],
        }

    @router.get("/summary")
    def summary(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        m = _assemble(request, domain, domains)
        states = [s for s in m["nat"].values() if not s["excluded"]]
        return {"as_of": m["now"].isoformat(), "scope": _scope_label(m), "dpo": m["dpo"], "blocked_by_dpo": not m["dpo"]["configured"],
                "gate": {k: m["gate"][k] for k in ("status", "rule")} | {"blockers": len(m["gate"]["blockers"]), "conditions": len(m["gate"]["conditions"])},
                "signoff": {k: m["sign"][k] for k in ("state", "signed", "total", "agreement")},
                "controls": {"total": len(states), "by_state": dict(Counter(s["state"] for s in states))},
                "risks": {"by_inherent": dict(Counter(r["inherent_band"] for r in m["risks"])),
                          "by_declared_residual": dict(Counter(r["residual_band"] for r in m["risks"])),
                          "by_evidence_adjusted": dict(Counter(r["adjusted_band"] for r in m["risks"]))},
                "axes": m["nat_axes"], "triangulation": m["tri"], "units": len(m["units"]),
                "workbook": {"source": m["wb_src"], "imported_at": m["wb"].get("imported_at"), "defects": len(m["wb"].get("defects", []))},
                "sources": m["sources"]}

    @router.get("/controls")
    def controls(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None),
                 axis: int | None = Query(None, ge=1, le=10), state: str | None = None, unit: str | None = None) -> dict:
        m = _assemble(request, domain, domains)
        src = m["nat"]
        if unit:
            u = m["units"].get(unit.lower())
            if not u:
                raise HTTPException(404, {"error": "unknown_unit"})
            if u["suppressed"]:
                raise HTTPException(422, {"error": "small_cell_suppressed", "k_min": m["cfg"]["k_min"]})
            src = u["_states"]
        rows = [s for s in src.values() if (axis is None or s["axis"] == axis) and (not state or s["state"] == state)]
        return {"as_of": m["now"].isoformat(), "unit": unit, "total": len(rows), "rows": rows,
                "axes": [{"axis": n, "name": a} for n, a in AXES]}

    @router.get("/risks")
    def risks(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        m = _assemble(request, domain, domains)
        return {"as_of": m["now"].isoformat(), "rows": m["risks"],
                "integrity": [{"risk": r["id"], "issue": r["flag"]} for r in m["risks"] if r["flag"]] +
                             [{"risk": r["id"], "issue": "no linked action"} for r in m["risks"] if not r["actions"]]}

    @router.get("/radar")
    def radar(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None), unit: str | None = None) -> dict:
        m = _assemble(request, domain, domains)
        sel = None
        if unit:
            u = m["units"].get(unit.lower())
            if not u:
                raise HTTPException(404, {"error": "unknown_unit"})
            sel = {"unit": u["unit"], "suppressed": u["suppressed"], "axes": m["unit_axes"][u["unit"]]}
        return {"as_of": m["now"].isoformat(), "scope": _scope_label(m), "axes": m["nat_axes"], "unit": sel,
                "series": ["ideal", "declared", "national_observed", "national_assured", "unit_observed", "threshold"],
                "units": [u["unit"] for u in _public_units(m)], "triangulation": m["tri"]}

    @router.get("/heatmap")
    def heatmap(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        m = _assemble(request, domain, domains)
        rows = []
        for u in _public_units(m):
            ax = m["unit_axes"][u["unit"]]
            rows.append({**u, "axes": [{"axis": a["axis"], "observed": a["observed"], "coverage": a["coverage"]} for a in ax]})
        rows.sort(key=lambda r: (r["overall"] is None, r["overall"] if r["overall"] is not None else 0))
        return {"as_of": m["now"].isoformat(), "k_min": m["cfg"]["k_min"], "axes": [{"axis": n, "name": a} for n, a in AXES], "rows": rows,
                "note": "Units are email domains. Scores are over the controls GAM can measure for that unit; many axes will be empty until attestations/collectors exist."}

    @router.get("/triangulation")
    def tri(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        m = _assemble(request, domain, domains)
        return {"as_of": m["now"].isoformat(), **m["tri"]}

    @router.get("/gate")
    def gate(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        m = _assemble(request, domain, domains)
        return {"as_of": m["now"].isoformat(), **m["gate"], "signoff_state": m["sign"]["state"]}

    @router.get("/workbook")
    def workbook_view(request: Request) -> dict:
        wb, src = _workbook()
        st = store.read_state()
        ov = st.get("tracker", {})
        now = _now()
        cfg = _cfg(st)
        acts = []
        for a in wb.get("actions", []):
            o = ov.get("action", {}).get(a["id"], {})
            tgt = o.get("target_date") or a["target"]
            stt = o.get("status") or a["status"]
            overdue = bool(tgt and (_parse_date(tgt) or now + timedelta(days=1)) < now and declared_value(stt) != 1.0)
            acts.append({**a, "status": stt, "owner": o.get("owner") or a["owner"], "target": tgt, "overdue": overdue,
                         "undated": not tgt, "evidence_ref": o.get("evidence_ref"), "note": o.get("note")})
        stk = []
        for s in wb.get("stakeholders", []):
            o = ov.get("stakeholder", {}).get(s["id"], {})
            stt = o.get("status") or s["status"]
            stk.append({**s, "status": stt, "evidence_ref": o.get("evidence_ref"), "note": o.get("note") or s["notes"],
                        "critical_path": s["required"].lower() == "yes" and declared_value(stt) == 0.0})
        return {"source": src, "imported_at": wb.get("imported_at"), "project": wb.get("project", {}), "actions": acts,
                "stakeholders": stk, "inventory": wb.get("inventory", []), "access": wb.get("access", []),
                "legal": wb.get("legal", []), "defects": wb.get("defects", []), "consult_sla_days": cfg["consult_sla_days"],
                "risk_to_control_map": RISK_CONTROLS}

    @router.get("/signoff")
    def signoff_get(request: Request) -> dict:
        wb, _ = _workbook()
        return signoff_view(wb, store.read_state().get("signoff") or {})

    @router.get("/settings")
    def settings_get(request: Request) -> dict:
        st = store.read_state()
        return {"settings": merged_settings(st.get("settings")), "dpo": _dpo(st), "defaults": DEFAULT_SETTINGS,
                "can_edit": bool(h.request_scope(request).get("superadmin"))}

    @router.get("/trend")
    def trend(request: Request, limit: int = Query(120, ge=1, le=1000)) -> dict:
        return {"rows": store.read_lines("snapshots.jsonl", limit)}

    @router.get("/report")
    def report(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> Response:
        m = _assemble(request, domain, domains)
        return Response(_report_md(m), media_type="text/markdown; charset=utf-8")

    @router.get("/feedback")
    def feedback_list(request: Request, limit: int = Query(200, ge=1, le=1000)) -> dict:
        _need_super(request)
        return {"rows": store.read_lines("feedback.jsonl", limit)}

    # ---------------- endpoints: write
    @router.put("/dpo")
    def set_dpo(body: DpoIn, request: Request) -> dict:
        _, who = _need_super(request)
        try:
            email, warns = validate_dpo_email(body.email)
        except ValueError as exc:
            raise HTTPException(422, {"error": "invalid_dpo_email", "message": str(exc)}) from exc
        rec = {"email": email, "name": (body.name or "").strip() or None, "role": (body.role or "").strip() or None,
               "designation_ref": (body.designation_ref or "").strip() or None, "set_by": who,
               "set_at": _now().isoformat(timespec="seconds")}
        store.mutate(lambda st: st.__setitem__("dpo", rec))
        _audit(who, "dpo_set", {"email": email})
        return {"ok": True, "dpo": {"configured": True, **rec}, "warnings": warns}

    @router.put("/settings")
    def settings_put(body: SettingsIn, request: Request) -> dict:
        _, who = _need_super(request)
        patch = body.model_dump(exclude_none=True)
        if "weights" in patch:
            bad = [k for k in patch["weights"] if k not in DEFAULT_SETTINGS["weights"] or not 0 <= patch["weights"][k] <= 1]
            if bad:
                raise HTTPException(422, {"error": "invalid_weights", "keys": bad, "allowed": list(DEFAULT_SETTINGS["weights"])})
        if "legal_status" in patch and patch["legal_status"].get("status") not in LEGAL_STATUSES:
            raise HTTPException(422, {"error": "invalid_legal_status", "allowed": list(LEGAL_STATUSES)})
        if "go_live_threshold" in patch:
            if any(k not in {str(n) for n, _ in AXES} or not 0 <= v <= 100 for k, v in patch["go_live_threshold"].items()):
                raise HTTPException(422, {"error": "invalid_threshold", "message": "keys '1'..'10', values 0-100"})
        if "excluded_controls" in patch and any(k not in CATALOGUE for k in patch["excluded_controls"]):
            raise HTTPException(422, {"error": "unknown_control"})
        st0 = store.read_state()
        _need_dpo(st0)

        def fn(st: dict[str, Any]) -> None:
            cur = st.setdefault("settings", {})
            for k, v in patch.items():
                if k == "weights":
                    cur.setdefault("weights", {}).update(v)
                else:
                    cur[k] = v
        store.mutate(fn)
        _audit(who, "settings_update", patch)
        return {"ok": True, "settings": merged_settings(store.read_state().get("settings"))}

    @router.post("/attestations")
    def attest(body: AttestationIn, request: Request) -> dict:
        sc, who = _actor(request)
        st0 = store.read_state()
        _need_dpo(st0)
        if body.control_id not in CATALOGUE:
            raise HTTPException(422, {"error": "unknown_control"})
        unit = (body.unit or "").strip().lower() or None
        if unit:
            if not sc.get("superadmin") and not h.in_scope(f"x@{unit}", set(sc.get("allowed_domains") or [])):
                raise HTTPException(403, {"error": "unit_out_of_scope"})
        elif not sc.get("superadmin"):
            raise HTTPException(403, {"error": "superadmin_required", "message": "national attestations are super-admin only"})
        if body.state == "partial" and body.pct is None:
            raise HTTPException(422, {"error": "pct_required", "message": "partial needs pct 0-1"})
        if body.valid_until and not _parse_date(body.valid_until):
            raise HTTPException(422, {"error": "invalid_valid_until"})
        rec = {**body.model_dump(), "unit": unit, "by": who, "at": _now().isoformat(timespec="seconds")}
        store.mutate(lambda st: st.setdefault("attestations", {}).__setitem__(f"{body.control_id}|{unit or ''}", rec))
        store.append("attestations.jsonl", rec)
        return {"ok": True, "attestation": rec,
                "tier": "evidenced" if body.evidence_ref.strip() else "attested",
                "note": "An attestation without an evidence link scores at the 'attested' weight (0.4 by default)."}

    @router.post("/tracker")
    def tracker(body: TrackerIn, request: Request) -> dict:
        _, who = _need_super(request)
        st0 = store.read_state()
        _need_dpo(st0)
        wb, _ = _workbook()
        ids = {a["id"] for a in wb.get("actions", [])} if body.kind == "action" else {s["id"] for s in wb.get("stakeholders", [])}
        if body.id not in ids:
            raise HTTPException(404, {"error": "unknown_id"})
        if body.target_date and not _parse_date(body.target_date):
            raise HTTPException(422, {"error": "invalid_target_date"})
        patch = {k: v for k, v in body.model_dump().items() if k not in ("kind", "id") and v is not None}
        patch.update(by=who, at=_now().isoformat(timespec="seconds"))
        store.mutate(lambda st: st.setdefault("tracker", {}).setdefault(body.kind, {}).setdefault(body.id, {}).update(patch))
        _audit(who, f"tracker_{body.kind}", {"id": body.id, **patch})
        return {"ok": True, "id": body.id, "patch": patch}

    @router.post("/signoff")
    def signoff_set(body: SignoffIn, request: Request) -> dict:
        _, who = _need_super(request)
        st0 = store.read_state()
        _need_dpo(st0)
        wb, _ = _workbook()
        if body.role and body.role not in (wb.get("signoff_roles") or []):
            raise HTTPException(404, {"error": "unknown_role", "roles": wb.get("signoff_roles")})
        if body.date and not _parse_date(body.date):
            raise HTTPException(422, {"error": "invalid_date"})

        def fn(st: dict[str, Any]) -> None:
            so = st.setdefault("signoff", {})
            if body.role:
                so.setdefault("roles", {})[body.role] = {"name": body.name or "", "date": body.date or "", "decision": body.decision or "",
                                                        "evidence_ref": body.evidence_ref or "", "by": who}
            if body.overall_decision:
                so["decision"] = {"value": body.overall_decision, "by": who, "at": _now().isoformat(timespec="seconds")}
            if body.conditions is not None:
                so["conditions"] = body.conditions
        store.mutate(fn)
        _audit(who, "signoff", body.model_dump(exclude_none=True))
        return {"ok": True, "signoff": signoff_view(wb, store.read_state().get("signoff") or {})}

    @router.post("/risk-acceptance")
    def accept_risk(body: AcceptanceIn, request: Request) -> dict:
        _, who = _need_super(request)
        _need_dpo(store.read_state())
        if not _parse_date(body.review_by):
            raise HTTPException(422, {"error": "invalid_review_by"})
        wb, _ = _workbook()
        if body.risk_id not in {r["id"] for r in wb.get("risks", [])}:
            raise HTTPException(404, {"error": "unknown_risk"})

        def fn(st: dict[str, Any]) -> None:
            acc = st.setdefault("acceptances", {})
            if body.revoke:
                acc.pop(body.risk_id, None)
            else:
                acc[body.risk_id] = {"note": body.note, "review_by": body.review_by, "by": who, "at": _now().isoformat(timespec="seconds")}
        store.mutate(fn)
        _audit(who, "risk_acceptance", body.model_dump())
        return {"ok": True}

    @router.post("/import")
    async def import_wb(request: Request) -> dict:
        _, who = _need_super(request)
        _need_dpo(store.read_state())
        raw = await request.body()
        if not raw or len(raw) > MAX_IMPORT_BYTES:
            raise HTTPException(413 if raw else 400, {"error": "bad_size", "max_bytes": MAX_IMPORT_BYTES})
        if raw[:2] != b"PK":
            raise HTTPException(415, {"error": "not_xlsx", "message": "send the .xlsx file as the raw request body"})
        try:
            parsed = parse_workbook(raw)
        except ValueError as exc:
            raise HTTPException(422, {"error": "unreadable_workbook", "message": str(exc)}) from exc
        if len(parsed["controls"]) < 20:
            raise HTTPException(422, {"error": "too_few_controls", "found": len(parsed["controls"])})
        ver = store.save_workbook(parsed)
        _audit(who, "workbook_import", {"version": ver, "controls": len(parsed["controls"]), "defects": len(parsed["defects"])})
        return {"ok": True, "version": ver, "controls": len(parsed["controls"]), "risks": len(parsed["risks"]),
                "actions": len(parsed["actions"]), "defects": parsed["defects"]}

    @router.post("/snapshot")
    def snapshot(request: Request, domain: list[str] | None = Query(None), domains: str | None = Query(None)) -> dict:
        _, who = _need_super(request)
        _need_dpo(store.read_state())
        m = _assemble(request, domain, domains)
        rec = {"at": m["now"].isoformat(timespec="seconds"), "by": who, "gate": m["gate"]["status"],
               "axes": {str(a["axis"]): {"observed": a["observed"], "assured": a["assured"], "declared": a["declared"]} for a in m["nat_axes"]},
               "coverage": m["tri"]["coverage"]}
        store.append("snapshots.jsonl", rec)
        return {"ok": True, "snapshot": rec}

    @router.post("/feedback")
    def feedback(body: FeedbackIn, request: Request) -> dict:
        _, who = _actor(request)
        rec = {**body.model_dump(), "by": who, "at": _now().isoformat(timespec="seconds"), "version": h.app_version}
        store.append("feedback.jsonl", rec)
        return {"ok": True, "message": "Thank you. Feedback is stored for the product owner."}

    return router


def _report_md(m: dict[str, Any]) -> str:
    d, g, t = m["dpo"], m["gate"], m["tri"]
    L = [f"# PIA compliance monitor report ({m['now']:%Y-%m-%d %H:%M} UTC)", "",
         f"Prepared for: **{d.get('name') or 'Data Protection Officer / focal point'}** <{d.get('email') or 'NOT CONFIGURED'}>", "",
         "_EXPERIMENTAL. Voluntary alignment with the Data Protection and Right to Access Information Bill 2025 (not enacted). "
         "Nothing below is a legal finding._", "", f"## Go-live gate: **{g['status']}**", ""]
    for b in g["blockers"]:
        L.append(f"- BLOCKER {b['ref']}: {b['message']}")
    for c in g["conditions"]:
        L.append(f"- CONDITION {c['ref']}: {c['message']}")
    L += ["", f"Sign-off: {m['sign']['state']} ({m['sign']['signed']}/{m['sign']['total']}). Agreement: {m['sign']['agreement']}.", "",
          "## Axes (0-100)", "", "| Axis | Declared | Observed | Assured | Coverage |", "|---|---|---|---|---|"]
    for a in m["nat_axes"]:
        L.append(f"| {a['axis']} {a['name']} | {a['declared']} | {a['observed'] if a['observed'] is not None else '–'} | {a['assured']} | {a['scored']}/{a['controls']} |")
    L += ["", f"Assurance gap: {t['assurance_gap']} · Ambition gap: {t['ambition_gap']} · Coverage: {t['coverage']:.0%}", "",
          "## Risks", "", "| Risk | Inherent | Declared residual | Evidence-adjusted |", "|---|---|---|---|"]
    for r in m["risks"]:
        L.append(f"| {r['id']} | {r['inherent']} {r['inherent_band']} | {r['residual']} {r['residual_band']} | {r['adjusted']} {r['adjusted_band']} |")
    L += ["", "## Workbook data-quality findings", ""]
    for x in m["wb"].get("defects", []):
        L.append(f"- [{x['severity']}] {x['where']}: {x['message']}")
    return "\n".join(L) + "\n"
