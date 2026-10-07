"""
Tests for app/privacy_monitor.py (0.4.20, experimental PIA compliance monitor).
No Postgres, no GAM, no network. Run from sl.p4sgi/apps/api:  python -m pytest tests -q
"""

from __future__ import annotations

import io
import json
import os
import shutil
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

_TMP = Path(tempfile.mkdtemp(prefix="p4sgi-priv-test-"))
os.environ["DATA_DIR"] = str(_TMP)
os.environ.setdefault("SUPERADMIN_EMAILS", "geb@p4sgi.com")
os.environ.setdefault("POSTGRES_HOST", "127.0.0.1")
os.environ.setdefault("POSTGRES_PORT", "1")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import privacy_monitor as pm  # noqa: E402
from app.privacy_catalogue import AXES, CONTROLS  # noqa: E402

NOW = datetime(2026, 10, 7, 12, 0, tzinfo=timezone.utc)
CFG = pm.merged_settings(None)
W = CFG["weights"]


def iso(days_ago: int) -> str:
    return (NOW - timedelta(days=days_ago)).strftime("%Y-%m-%dT%H:%M:%S.000Z")


# ------------------------------------------------------------ catalogue/seed
def test_catalogue_is_40_controls_10_axes_and_matches_seed():
    ids = [c[0] for c in CONTROLS]
    assert ids == [f"A{i:02d}" for i in range(1, 41)]
    assert [a[0] for a in AXES] == list(range(1, 11))
    assert {c[1] for c in CONTROLS} == set(range(1, 11))
    seed = pm.load_seed()
    assert [c["id"] for c in seed["controls"]] == ids
    assert len(seed["risks"]) == 18 and len(seed["actions"]) == 16 and len(seed["stakeholders"]) == 13
    assert len(seed["signoff_roles"]) == 6
    assert all(a["action"] for a in seed["actions"])  # action text column picked up, not the id column
    assert set(pm.RISK_CONTROLS) == {f"R{i}" for i in range(1, 19)}
    assert all(c in pm.CATALOGUE for v in pm.RISK_CONTROLS.values() for c in v)


def test_seed_defects_surface_the_known_workbook_problems():
    msgs = " | ".join(d["message"] for d in pm.load_seed()["defects"])
    assert "mitigation credited before it exists" in msgs
    assert "retention is approved for 0 of 16" in msgs
    assert "R4 has no linked mitigation action" in msgs


def test_bands_follow_the_scoring_guide():
    assert [pm.band(x) for x in (1, 4, 5, 9, 10, 15, 16, 25)] == ["Low", "Low", "Medium", "Medium", "High", "High", "Critical", "Critical"]
    assert pm.band(None) is None


def test_declared_vocabulary_unknown_is_none():
    assert pm.declared_value("Not started") == 0.0 and pm.declared_value("In progress") == 0.5
    assert pm.declared_value("Completed") == 1.0 and pm.declared_value("whatever") is None


# --------------------------------------------------------------- DPO email
@pytest.mark.parametrize("bad", ["", "   ", "nope", "a@b", "test@company.org", "dpo@example.com", "x@localhost", "a b@c.org"])
def test_dpo_email_rejects_invalid_and_placeholders(bad):
    with pytest.raises(ValueError):
        pm.validate_dpo_email(bad)


def test_dpo_email_normalises_and_warns_on_personal_mailbox():
    e, w = pm.validate_dpo_email("  Focal.Point@MBSSE.gov.sl ")
    assert e == "focal.point@mbsse.gov.sl" and w == []
    e, w = pm.validate_dpo_email("someone@gmail.com")
    assert w and w[0].startswith("personal_mailbox")


# ----------------------------------------------------------------- scoring
def test_attestation_tiers_and_partial_cap():
    a = {"state": "met", "evidence_ref": "", "at": NOW.isoformat()}
    assert pm.attestation_result(a, W, CFG, NOW)["score"] == 0.4
    a["evidence_ref"] = "vault://matter/123"
    assert pm.attestation_result(a, W, CFG, NOW)["score"] == 0.75
    p = {"state": "partial", "pct": 1.0, "evidence_ref": "doc", "at": NOW.isoformat()}
    assert pm.attestation_result(p, W, CFG, NOW)["score"] == round(0.9 * 0.75, 4)  # partial is capped at 0.9
    assert pm.attestation_result({"state": "gap", "evidence_ref": "doc", "at": NOW.isoformat()}, W, CFG, NOW)["score"] == 0.0


def test_expired_evidence_counts_half_then_zero():
    base = {"state": "met", "evidence_ref": "doc"}
    half = pm.attestation_result({**base, "valid_until": (NOW - timedelta(days=10)).isoformat()}, W, CFG, NOW)
    assert half["expired"] and half["score"] == 0.375 and not half["void"]
    void = pm.attestation_result({**base, "valid_until": (NOW - timedelta(days=400)).isoformat()}, W, CFG, NOW)
    assert void["void"] and void["score"] == 0.0


def test_collector_never_reaches_verified_without_evidenced_attestation():
    assert pm.collector_result(1.0, 0.9, W, CFG) == ("partial", 0.9)
    assert pm.collector_result(1.0, 1.0, W, CFG) == ("verified", 1.0)
    assert pm.collector_result(0.0, 1.0, W, CFG)[0] == "gap"
    assert pm.collector_result(0.5, 1.0, W, CFG) == ("partial", 0.5)


def _states(measured=None, attests=None, dpo=None, cfg=CFG, unit=None):
    return pm.control_states(pm.load_seed(), measured or {}, attests or {}, cfg, unit, dpo, NOW)


def test_nothing_is_observed_from_the_workbook_alone():
    s = _states()
    assert all(v["score"] is None and v["state"] == "no_data" for v in s.values())
    axes = pm.axis_scores(s, CFG)
    assert all(a["observed"] is None and a["coverage"] == 0 and a["assured"] == 0 and a["declared"] == 0 for a in axes)


def test_dpo_setting_gives_partial_attested_credit_to_a02_only():
    s = _states(dpo="dpo@mbsse.gov.sl")
    assert s["A02"]["state"] == "partial" and s["A02"]["basis"] == "dpo_setting" and s["A02"]["score"] == 0.2
    assert [k for k, v in s.items() if v["score"] is not None] == ["A02"]
    assert _states(dpo="x@y.org", unit="sl.p4sgi.com")["A02"]["score"] is None  # national control, not copied to units


def test_attestation_conflicts_with_contrary_measurement_and_measurement_wins():
    att = {"A22|": {"control_id": "A22", "state": "met", "evidence_ref": "doc", "at": NOW.isoformat(), "by": "x"}}
    s = _states({"A22": {"frac": 0.4, "detail": "2SV 40%", "n": 10}}, att)
    assert s["A22"]["conflict"] and s["A22"]["score"] == 0.4 and s["A22"]["basis"] == "measured"


def test_evidenced_attestation_lifts_the_collector_cap():
    m = {"A22": {"frac": 1.0, "detail": "all", "n": 10}}
    assert _states(m)["A22"]["state"] == "partial" and _states(m)["A22"]["score"] == 0.9
    att = {"A22|": {"control_id": "A22", "state": "met", "evidence_ref": "vault://x", "at": NOW.isoformat(), "by": "x"}}
    assert _states(m, att)["A22"]["state"] == "verified" and _states(m, att)["A22"]["score"] == 1.0
    bare = {"A22|": {"control_id": "A22", "state": "met", "evidence_ref": "", "at": NOW.isoformat(), "by": "x"}}
    assert _states(m, bare)["A22"]["score"] == 0.9


def test_unit_attestation_does_not_leak_to_national_or_other_units():
    att = {"A06|sl.p4sgi.com": {"control_id": "A06", "state": "met", "evidence_ref": "d", "at": NOW.isoformat(), "by": "x"}}
    assert _states(attests=att, unit="sl.p4sgi.com")["A06"]["score"] == 0.75
    assert _states(attests=att, unit=None)["A06"]["score"] is None
    assert _states(attests=att, unit="zw.p4sgi.com")["A06"]["score"] is None


def test_excluded_control_is_left_out_of_scoring_with_reason():
    cfg = pm.merged_settings({"excluded_controls": {"A15": "no biometrics in scope"}})
    s = _states(cfg=cfg)
    assert s["A15"]["excluded"] and "no biometrics" in s["A15"]["detail"]
    ax = pm.axis_scores(s, cfg)
    assert next(a for a in ax if a["axis"] == 5)["controls"] == 3


def test_computed_controls():
    s = _states()
    pm.finalize_computed(s, False, NOW)
    assert s["A40"]["state"] == "gap" and s["A37"]["state"] == "gap"
    pm.finalize_computed(s, True, NOW)
    assert s["A40"]["state"] == "verified"


# ------------------------------------------------------------ risks and gate
def _risks(states, overlay=None, acc=None):
    return pm.risk_view(pm.load_seed(), states, overlay or {}, acc or {}, CFG, NOW)


def test_residual_risk_is_not_credited_without_completed_actions_and_evidence():
    r = {x["id"]: x for x in _risks(_states())}
    assert r["R1"]["inherent"] == 20 and r["R1"]["residual"] == 12
    assert r["R1"]["adjusted"] == 20 and r["R1"]["adjusted_band"] == "Critical" and not r["R1"]["credited"]
    assert r["R4"]["why_not_credited"][0] == "no linked action"
    assert r["R15"]["flag"] == "accepted without an acceptance record"


def test_residual_credit_needs_both_complete_actions_and_evidence_in_tolerance():
    att = {f"{c}|": {"control_id": c, "state": "met", "evidence_ref": "doc", "at": NOW.isoformat(), "by": "x"} for c in ("A02",)}
    s = _states(attests=att)
    done = {"ACT-01": {"status": "Completed"}, "ACT-14": {"status": "Completed"}}  # R16 is linked to both
    r16 = next(x for x in _risks(s, done) if x["id"] == "R16")
    assert r16["credited"] and r16["adjusted"] == r16["residual"] == 20  # workbook residual for R16 is unchanged (5x4)
    only_actions = next(x for x in _risks(_states(), done) if x["id"] == "R16")
    assert not only_actions["credited"]
    only_evidence = next(x for x in _risks(s) if x["id"] == "R16")
    assert not only_evidence["credited"]


def test_gate_blocks_without_dpo_and_open_only_when_everything_is_clear():
    seed = pm.load_seed()
    sign_ok = {"state": "Signed", "signed": 6, "total": 6}
    sign_no = {"state": "Draft", "signed": 0, "total": 6}
    g = pm.gate_view(_risks(_states()), sign_no, None)
    assert g["status"] == "BLOCKED" and {b["kind"] for b in g["blockers"]} >= {"dpo_missing", "signoff", "risk_critical"}
    acc = {r["id"]: {"review_by": (NOW + timedelta(days=30)).isoformat()} for r in seed["risks"]}
    g2 = pm.gate_view(_risks(_states(), acc=acc), sign_ok, "dpo@mbsse.gov.sl")
    assert g2["status"] == "OPEN"
    # expired acceptance stops counting
    lapsed = {r["id"]: {"review_by": (NOW - timedelta(days=1)).isoformat()} for r in seed["risks"]}
    assert pm.gate_view(_risks(_states(), acc=lapsed), sign_ok, "d@mbsse.gov.sl")["status"] == "BLOCKED"


def test_gate_conditional_when_only_high_risks_remain():
    risks = [{"id": "R9", "adjusted_band": "High", "adjusted": 12, "accepted": False, "why_not_credited": []}]
    g = pm.gate_view(risks, {"state": "Signed", "signed": 6, "total": 6}, "d@mbsse.gov.sl")
    assert g["status"] == "CONDITIONAL" and not g["blockers"]


def test_signoff_state_machine():
    seed = pm.load_seed()
    roles = seed["signoff_roles"]
    assert pm.signoff_view(seed, {})["state"] == "Draft" and pm.signoff_view(seed, {})["agreement"] == "not evidenced"
    one = {"roles": {roles[0]: {"name": "A", "date": "2026-10-01"}}}
    assert pm.signoff_view(seed, one)["state"] == "In review"
    allr = {"roles": {r: {"name": "A", "date": "2026-10-01"} for r in roles}}
    assert pm.signoff_view(seed, allr)["state"] == "Signed"
    allr["decision"] = {"value": "Proceed with conditions"}
    v = pm.signoff_view(seed, allr)
    assert v["state"] == "Decision recorded" and v["agreement"] == "evidenced"
    assert pm.signoff_view(seed, {"roles": {roles[0]: {"name": "A"}}})["state"] == "In review"  # name without date is not signed


# --------------------------------------------------------------- collectors
def _em(row, *cols):
    for c in cols:
        v = (row.get(c) or "").strip().lower()
        if "@" in v:
            return v
    return ""


def test_compute_collectors_aggregates_only():
    users = [{"primaryEmail": f"u{i}@sl.p4sgi.com", "suspended": "False", "isEnrolledIn2Sv": "True" if i < 3 else "False",
              "lastLoginTime": iso(2), "creationTime": iso(300)} for i in range(4)]
    users.append({"primaryEmail": "gone@sl.p4sgi.com", "suspended": "True", "isEnrolledIn2Sv": "False", "lastLoginTime": iso(2)})
    admins = [{"primaryEmail": "u0@sl.p4sgi.com"}, {"primaryEmail": "u3@sl.p4sgi.com"}]
    meas, n = pm.compute_collectors(users, [], admins, [], {"login_activity": 5}, _em, NOW)
    assert n == 4 and meas["A22"]["frac"] == 0.75 and meas["A20"]["frac"] == 0.5
    assert meas["A23"]["frac"] == pytest.approx(1 / 3)
    none_cached, _ = pm.compute_collectors(users, [], admins, [], {"login_activity": None, "token_activity": None, "drive_activity": None}, _em, NOW)
    assert "A23" not in none_cached  # uncached audit reports are 'no data', never a gap
    assert "u0@" not in json.dumps(meas)  # no person-level data in the output
    assert "A26" not in meas  # no devices -> no device claims


def test_compute_collectors_without_flag_columns_makes_no_claim():
    users = [{"primaryEmail": "a@sl.p4sgi.com", "lastLoginTime": iso(1)}]
    meas, _ = pm.compute_collectors(users, [], [], [], {}, _em, NOW)
    assert "A22" not in meas and "A20" not in meas


# ------------------------------------------------------------------- import
def _build_xlsx(controls: int = 40, bad_score: bool = False) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Instructions"
    ws = wb.create_sheet("3. Country Assessment")
    ws.append(["Country / Project Assessment"])
    ws.append(["ID", "Assessment area", "Control / question for country validation", "Status", "Evidence / reference",
               "Gap / finding", "Corrective action", "Owner", "Target date"])
    for i in range(1, controls + 1):
        ws.append([f"A{i:02d}", "Governance", f"Question {i}", "In progress" if i == 1 else "Not started"])
    ws = wb.create_sheet("4. Risk Register")
    ws.append(["ID", "Risk Description", "Category", "Data Subjects Affected", "Likelihood\n(1-5)", "Impact\n(1-5)", "Inherent\nScore",
               "Inherent\nRating", "Existing / Planned Mitigation", "Residual\nLikelihood", "Residual\nImpact", "Residual\nScore",
               "Residual\nRating", "Risk Owner", "Status"])
    ws.append(["R1", "No law", "Legal", "All", 5, 4, 19 if bad_score else 20, "Critical", "x", 6 if bad_score else 4, 5, 20, "Critical", "MBSSE", "Open"])
    ws = wb.create_sheet("6. Mitigation Tracker")
    ws.append(["Action ID", "Action", "Linked Risk ID(s)", "Owner", "Priority", "Target Date", "Status"])
    ws.append(["ACT-01", "Designate DPO", "R1, R99", "MBSSE", "Critical", "[Enter date]", "Not started"])
    ws = wb.create_sheet("10. Sign-off")
    ws.append(["Role", "Name", "Signature", "Date", "Decision / Comments"])
    ws.append(["Approved by - PS"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_workbook_reads_tabs_and_reports_defects_without_fixing_them():
    d = pm.parse_workbook(_build_xlsx(bad_score=True))
    assert len(d["controls"]) == 40 and d["controls"][0]["declared"] == 0.5
    assert d["actions"][0]["action"] == "Designate DPO" and d["actions"][0]["risks"] == ["R1", "R99"]
    msgs = " | ".join(x["message"] for x in d["defects"])
    assert "inherent score 19 != L x I = 20" in msgs and "residual 30 is higher than inherent 20" in msgs.replace("residual 30", "residual 30") or "higher than inherent" in msgs
    assert "unknown risk R99" in msgs
    assert d["risks"][0]["inherent"] == 20  # recomputed, not the typed 19


def test_parse_workbook_rejects_non_pia_and_garbage():
    with pytest.raises(ValueError):
        pm.parse_workbook(b"not a zip")
    import openpyxl

    buf = io.BytesIO()
    openpyxl.Workbook().save(buf)
    with pytest.raises(ValueError):
        pm.parse_workbook(buf.getvalue())


# --------------------------------------------------------- HTTP integration
def _priv() -> Path:
    from app import main

    return Path(main.DATA_DIR) / "privacy"


def _write_run(report: str, csv_text: str, stamp: str = "20261007T080000Z") -> None:
    runs = _TMP / "gam-out" / "runs"
    runs.mkdir(parents=True, exist_ok=True)
    (runs / f"{stamp}_{report}.csv").write_text(csv_text, encoding="utf-8")
    (runs / f"{stamp}_{report}.meta.json").write_text(json.dumps(
        {"report": report, "status": "ok", "created_at": "2026-10-07T08:00:00Z", "domains": [], "filename": f"{stamp}_{report}.csv"}), encoding="utf-8")


@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    _write_run("users_full",
               "primaryEmail,name.fullName,orgUnitPath,creationTime,lastLoginTime,suspended,isEnrolledIn2Sv\n"
               f"a@sl.p4sgi.com,A,/SL,{iso(300)},{iso(2)},False,True\n"
               f"a2@sl.p4sgi.com,A2,/SL,{iso(300)},{iso(2)},False,False\n"
               f"b@zw.p4sgi.com,B,/ZW,{iso(300)},{iso(2)},False,False\n")
    _write_run("devices_mobile",
               "deviceId,email,model,os,type,status,lastSync,encryptionStatus,developerOptionsStatus,adbStatus,"
               "unknownSourcesStatus,devicePasswordStatus,securityPatchLevel,serialNumber\n"
               f"d1,a@sl.p4sgi.com,SM-X216B,Android 14,ANDROID,APPROVED,{iso(1)},NOT_ENCRYPTED,true,false,false,PASSWORD_SET,2024-01-01,SER1\n")
    from app import main

    # Other test modules share this process and the app's DATA_DIR: give this module its own GAM run cache.
    saved = main.GAM_RUNS_DIR
    main.GAM_RUNS_DIR = _TMP / "gam-out" / "runs"
    yield TestClient(main.app)
    main.GAM_RUNS_DIR = saved


SL = {"X-Goog-Authenticated-User-Email": "accounts.google.com:ht@sl.p4sgi.com"}
SUPER = {"X-Goog-Authenticated-User-Email": "accounts.google.com:geb@p4sgi.com"}
P = "/api/v1/privacy"


def _reset():
    shutil.rmtree(_priv(), ignore_errors=True)


def _set_dpo(client, email="focal.point@mbsse.gov.sl"):
    return client.put(f"{P}/dpo", json={"email": email, "name": "Focal Point"}, headers=SUPER)


def test_status_is_experimental_and_honest_about_limits(client):
    _reset()
    j = client.get(f"{P}/status", headers=SL).json()
    assert j["experimental"] is True and j["dpo"]["configured"] is False and j["workbook_source"] == "bundled_seed"
    assert j["counts"]["controls"] == 40 and "NOT enacted" in j["legal"]["framework"]
    assert any("7 of 40" in s for s in j["limits"])


def test_dpo_is_required_before_any_write_except_dpo_and_feedback(client):
    _reset()
    s = client.get(f"{P}/summary", headers=SUPER).json()
    assert s["blocked_by_dpo"] is True and s["gate"]["status"] == "BLOCKED"
    body = {"control_id": "A06", "state": "met", "evidence_ref": "doc"}
    r = client.post(f"{P}/attestations", json=body, headers=SUPER)
    assert r.status_code == 409 and r.json()["detail"]["error"] == "dpo_required"
    for path, payload in [("/tracker", {"kind": "action", "id": "ACT-01", "status": "In progress"}),
                          ("/signoff", {"conditions": "x"}), ("/snapshot", None)]:
        r = client.post(f"{P}{path}", json=payload, headers=SUPER) if payload is not None else client.post(f"{P}{path}", headers=SUPER)
        assert r.status_code == 409, path
    assert client.put(f"{P}/settings", json={"k_min": 5}, headers=SUPER).status_code == 409
    assert client.post(f"{P}/feedback", json={"text": "looks useful", "category": "idea"}, headers=SL).status_code == 200


def test_set_dpo_validation_and_permissions(client):
    _reset()
    assert client.put(f"{P}/dpo", json={"email": "dpo@mbsse.gov.sl"}, headers=SL).status_code == 403
    r = client.put(f"{P}/dpo", json={"email": "not-an-email"}, headers=SUPER)
    assert r.status_code == 422 and r.json()["detail"]["error"] == "invalid_dpo_email"
    assert client.put(f"{P}/dpo", json={"email": "dpo@example.com"}, headers=SUPER).status_code == 422
    r = _set_dpo(client)
    assert r.status_code == 200 and r.json()["dpo"]["email"] == "focal.point@mbsse.gov.sl"
    st = client.get(f"{P}/status", headers=SL).json()
    assert st["dpo"]["configured"] is True and st["dpo"]["email"] == "focal.point@mbsse.gov.sl"
    audit = (_priv() / "audit.jsonl").read_text()
    assert "dpo_set" in audit


def test_summary_measures_gam_controls_and_gives_nothing_else_credit(client):
    _reset()
    _set_dpo(client)
    s = client.get(f"{P}/summary", headers=SUPER).json()
    by = s["controls"]["by_state"]
    assert by["partial"] >= 3 and by["no_data"] >= 30 and "verified" not in by
    assert s["gate"]["status"] == "BLOCKED"
    c = {x["id"]: x for x in client.get(f"{P}/controls", headers=SUPER).json()["rows"]}
    assert c["A22"]["basis"] == "measured" and round(c["A22"]["measured"]["frac"], 2) == 0.33
    assert c["A26"]["measured"]["frac"] == 0.0 and c["A26"]["state"] == "gap"
    assert c["A02"]["basis"] == "dpo_setting"
    assert c["A01"]["state"] == "no_data"
    assert s["triangulation"]["coverage"] < 0.3
    assert s["signoff"]["agreement"] == "not evidenced"


def test_domain_scoping_for_non_superadmin(client):
    _reset()
    _set_dpo(client)
    own = client.get(f"{P}/heatmap", headers=SL).json()
    assert [u["unit"] for u in own["rows"]] == ["sl.p4sgi.com"]
    allr = client.get(f"{P}/heatmap", headers=SUPER).json()
    assert {u["unit"] for u in allr["rows"]} == {"sl.p4sgi.com", "zw.p4sgi.com"}


def test_small_cells_are_suppressed_until_k_min_is_lowered(client):
    _reset()
    _set_dpo(client)
    h = client.get(f"{P}/heatmap", headers=SUPER).json()
    assert all(u["suppressed"] and u["overall"] is None and u["accounts"] is None for u in h["rows"])
    assert client.get(f"{P}/controls?unit=sl.p4sgi.com", headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"k_min": 1}, headers=SUPER).status_code == 200
    h = client.get(f"{P}/heatmap", headers=SUPER).json()
    sl = next(u for u in h["rows"] if u["unit"] == "sl.p4sgi.com")
    assert not sl["suppressed"] and sl["accounts"] == 2 and sl["overall"] is not None
    rd = client.get(f"{P}/radar?unit=sl.p4sgi.com", headers=SUPER).json()
    assert rd["unit"]["unit"] == "sl.p4sgi.com" and len(rd["axes"]) == 10
    assert client.get(f"{P}/radar?unit=nope.org", headers=SUPER).status_code == 404


def test_attestations_permissions_and_effect(client):
    _reset()
    _set_dpo(client)
    body = {"control_id": "A06", "state": "met", "evidence_ref": "https://drive/x", "note": "notice on enrolment screen"}
    assert client.post(f"{P}/attestations", json=body, headers=SL).status_code == 403          # national needs super-admin
    own = client.post(f"{P}/attestations", json={**body, "unit": "sl.p4sgi.com"}, headers=SL)
    assert own.status_code == 200 and own.json()["tier"] == "evidenced"
    assert client.post(f"{P}/attestations", json={**body, "unit": "zw.p4sgi.com"}, headers=SL).status_code == 403
    assert client.post(f"{P}/attestations", json={"control_id": "A99", "state": "met"}, headers=SUPER).status_code == 422
    assert client.post(f"{P}/attestations", json={"control_id": "A06", "state": "partial"}, headers=SUPER).status_code == 422
    assert client.post(f"{P}/attestations", json=body, headers=SUPER).status_code == 200
    c = {x["id"]: x for x in client.get(f"{P}/controls", headers=SUPER).json()["rows"]}
    assert c["A06"]["state"] == "evidenced" and c["A06"]["score"] == 0.75
    assert "A06|" in json.loads((_priv() / "state.json").read_text())["attestations"]
    assert (_priv() / "attestations.jsonl").exists()


def test_risks_gate_signoff_and_acceptance_flow(client):
    _reset()
    _set_dpo(client)
    rs = client.get(f"{P}/risks", headers=SUPER).json()
    assert len(rs["rows"]) == 18
    assert {i["risk"] for i in rs["integrity"]} >= {"R4", "R15", "R17"}
    g = client.get(f"{P}/gate", headers=SUPER).json()
    assert g["status"] == "BLOCKED" and not any(b["kind"] == "dpo_missing" for b in g["blockers"])
    # sign-off
    roles = client.get(f"{P}/signoff", headers=SUPER).json()["rows"]
    assert client.post(f"{P}/signoff", json={"role": "nope", "name": "x"}, headers=SUPER).status_code == 404
    for r in roles:
        assert client.post(f"{P}/signoff", json={"role": r["role"], "name": "Signatory", "date": "2026-10-01"}, headers=SUPER).status_code == 200
    so = client.post(f"{P}/signoff", json={"overall_decision": "Proceed with conditions", "conditions": "close ACT-01"}, headers=SUPER).json()["signoff"]
    assert so["state"] == "Decision recorded" and so["agreement"] == "evidenced"
    # accept every Critical/High risk with a future review date
    for r in rs["rows"]:
        a = client.post(f"{P}/risk-acceptance", json={"risk_id": r["id"], "note": "accepted for pilot", "review_by": "2027-06-01"}, headers=SUPER)
        assert a.status_code == 200
    g = client.get(f"{P}/gate", headers=SUPER).json()
    assert g["status"] == "OPEN"
    client.post(f"{P}/risk-acceptance", json={"risk_id": "R1", "note": "revoke", "review_by": "2027-06-01", "revoke": True}, headers=SUPER)
    assert client.get(f"{P}/gate", headers=SUPER).json()["status"] == "BLOCKED"


def test_tracker_overlay_marks_overdue_and_credits_residual_only_with_evidence(client):
    _reset()
    _set_dpo(client)
    assert client.post(f"{P}/tracker", json={"kind": "action", "id": "ACT-99", "status": "x"}, headers=SUPER).status_code == 404
    assert client.post(f"{P}/tracker", json={"kind": "action", "id": "ACT-01", "target_date": "garbage"}, headers=SUPER).status_code == 422
    assert client.post(f"{P}/tracker", json={"kind": "action", "id": "ACT-02", "target_date": "2026-01-01"}, headers=SUPER).status_code == 200
    wbv = client.get(f"{P}/workbook", headers=SUPER).json()
    a2 = next(a for a in wbv["actions"] if a["id"] == "ACT-02")
    assert a2["overdue"] is True and next(a for a in wbv["actions"] if a["id"] == "ACT-03")["undated"] is True
    client.post(f"{P}/tracker", json={"kind": "action", "id": "ACT-01", "status": "Completed"}, headers=SUPER)
    client.post(f"{P}/tracker", json={"kind": "action", "id": "ACT-14", "status": "Completed"}, headers=SUPER)
    r16 = next(r for r in client.get(f"{P}/risks", headers=SUPER).json()["rows"] if r["id"] == "R16")
    assert r16["credited"] is False  # action says Completed, but the evidence for A02 is below tolerance
    client.post(f"{P}/attestations", json={"control_id": "A02", "state": "met", "evidence_ref": "vault://designation-letter"}, headers=SUPER)
    r16 = next(r for r in client.get(f"{P}/risks", headers=SUPER).json()["rows"] if r["id"] == "R16")
    assert r16["credited"] is True


def test_settings_validation(client):
    _reset()
    _set_dpo(client)
    assert client.put(f"{P}/settings", json={"k_min": 0}, headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"weights": {"bogus": 0.5}}, headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"legal_status": {"status": "Enacted"}}, headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"go_live_threshold": {"11": 50}}, headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"excluded_controls": {"A99": "x"}}, headers=SUPER).status_code == 422
    assert client.put(f"{P}/settings", json={"k_min": 5}, headers=SL).status_code == 403
    ok = client.put(f"{P}/settings", json={"weights": {"attested": 0.3}, "go_live_threshold": {"7": 80},
                                           "legal_status": {"status": "Assented", "date": "2026-12-01"}}, headers=SUPER)
    assert ok.status_code == 200 and ok.json()["settings"]["weights"]["attested"] == 0.3
    ax = client.get(f"{P}/radar", headers=SUPER).json()["axes"]
    assert next(a for a in ax if a["axis"] == 7)["threshold"] == 80
    assert client.get(f"{P}/status", headers=SL).json()["legal"]["status"] == "Assented"


def test_import_replaces_seed_and_rejects_bad_files(client):
    _reset()
    assert client.post(f"{P}/import", content=_build_xlsx(), headers=SUPER).status_code == 409   # DPO first
    _set_dpo(client)
    assert client.post(f"{P}/import", content=b"hello", headers=SUPER).status_code == 415
    assert client.post(f"{P}/import", content=b"", headers=SUPER).status_code == 400
    assert client.post(f"{P}/import", content=_build_xlsx(), headers=SL).status_code == 403
    assert client.post(f"{P}/import", content=_build_xlsx(controls=5), headers=SUPER).status_code == 422
    r = client.post(f"{P}/import", content=_build_xlsx(bad_score=True), headers=SUPER)
    assert r.status_code == 200 and r.json()["controls"] == 40 and r.json()["defects"]
    assert client.get(f"{P}/status", headers=SL).json()["workbook_source"] == "imported"
    assert list((_priv() / "workbooks").glob("*.json"))  # versioned copy kept
    assert len(client.get(f"{P}/risks", headers=SUPER).json()["rows"]) == 1


def test_snapshot_trend_report_and_feedback(client):
    _reset()
    _set_dpo(client)
    assert client.post(f"{P}/snapshot", headers=SL).status_code == 403
    assert client.post(f"{P}/snapshot", headers=SUPER).status_code == 200
    assert len(client.get(f"{P}/trend", headers=SL).json()["rows"]) == 1
    md = client.get(f"{P}/report", headers=SUPER)
    assert md.status_code == 200 and "focal.point@mbsse.gov.sl" in md.text and "Go-live gate: **BLOCKED**" in md.text
    assert "not enacted" in md.text
    assert client.post(f"{P}/feedback", json={"text": "x"}, headers=SL).status_code == 422
    assert client.post(f"{P}/feedback", json={"text": "A22 should count staff only", "category": "wrong_control", "control_id": "A22"}, headers=SL).status_code == 200
    assert client.post(f"{P}/feedback", json={"text": "bad id", "control_id": "ZZ"}, headers=SL).status_code == 422
    assert client.get(f"{P}/feedback", headers=SL).status_code == 403
    fb = client.get(f"{P}/feedback", headers=SUPER).json()["rows"]
    assert fb[-1]["control_id"] == "A22" and fb[-1]["by"] == "ht@sl.p4sgi.com"


def test_no_learner_level_data_is_stored(client):
    _reset()
    _set_dpo(client)
    client.get(f"{P}/summary", headers=SUPER)
    client.post(f"{P}/snapshot", headers=SUPER)
    blob = "".join(p.read_text() for p in (_priv()).rglob("*") if p.is_file())
    assert "a2@sl.p4sgi.com" not in blob and "SER1" not in blob
